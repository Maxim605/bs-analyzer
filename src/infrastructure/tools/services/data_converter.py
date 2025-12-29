from __future__ import annotations
from typing import Dict, Any, List
from io import BytesIO
import numpy as np
from openpyxl import Workbook, load_workbook

from src.domain.tools.services.converter_interface import IDataConverter


class DataConverterService(IDataConverter):
    """Реализация конвертера данных."""

    def d3_json_to_adjacency_matrix(
        self, 
        nodes: List[Dict[str, Any]], 
        links: List[Dict[str, Any]]
    ) -> tuple[np.ndarray, List[str]]:
        """Конвертирует D3 JSON граф в матрицу смежности."""
        # Создаем маппинг id -> index
        node_ids = [str(node.get("id", i)) for i, node in enumerate(nodes)]
        id_to_idx = {node_id: idx for idx, node_id in enumerate(node_ids)}
        
        n = len(nodes)
        matrix = np.zeros((n, n), dtype=np.float64)
        
        # Заполняем матрицу связями
        for link in links:
            source_raw = link.get("source", "")
            target_raw = link.get("target", "")
            
            # source/target могут быть объектами или строками
            if isinstance(source_raw, dict):
                source = str(source_raw.get("id", ""))
            else:
                source = str(source_raw)
            
            if isinstance(target_raw, dict):
                target = str(target_raw.get("id", ""))
            else:
                target = str(target_raw)
            
            weight_val = link.get("weight")
            weight = float(weight_val) if weight_val is not None else 1.0
            
            if source in id_to_idx and target in id_to_idx:
                src_idx = id_to_idx[source]
                tgt_idx = id_to_idx[target]
                matrix[src_idx, tgt_idx] = weight
                # Для неориентированного графа добавляем симметричную связь
                matrix[tgt_idx, src_idx] = weight
        
        return matrix, node_ids

    def adjacency_matrix_to_d3_json(
        self, 
        matrix: np.ndarray, 
        node_ids: List[str]
    ) -> Dict[str, Any]:
        """Конвертирует матрицу смежности в D3 JSON формат."""
        n = matrix.shape[0]
        
        # Создаем узлы - используем id из заголовков
        nodes = []
        for i, node_id in enumerate(node_ids):
            nodes.append({
                "id": node_id,
                "name": node_id,
                "index": i
            })
        
        # Создаем связи (только верхний треугольник для неориентированного графа)
        links = []
        for i in range(n):
            for j in range(i + 1, n):
                weight = matrix[i, j]
                if weight != 0:
                    link = {
                        "source": node_ids[i],
                        "target": node_ids[j]
                    }
                    if weight != 1.0:
                        link["weight"] = float(weight)
                    links.append(link)
        
        return {"nodes": nodes, "links": links}

    def matrix_to_xlsx_bytes(
        self, 
        matrix: np.ndarray, 
        headers: List[str]
    ) -> bytes:
        """Конвертирует матрицу в XLSX формат."""
        wb = Workbook(write_only=True)
        ws = wb.create_sheet(title="Adjacency Matrix")
        
        # Первая строка - заголовки
        header_row = [""] + headers
        ws.append(header_row)
        
        # Записываем данные матрицы построчно
        for i in range(matrix.shape[0]):
            row_header = headers[i] if i < len(headers) else str(i)
            row_data = [row_header]
            for j in range(matrix.shape[1]):
                value = matrix[i, j]
                # Записываем как int если целое, иначе float
                if value == int(value):
                    row_data.append(int(value))
                else:
                    row_data.append(float(value))
            ws.append(row_data)
        
        # Сохраняем в байты
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    def xlsx_bytes_to_matrix(
        self, 
        xlsx_bytes: bytes
    ) -> tuple[np.ndarray, List[str]]:
        """Читает матрицу из XLSX файла."""
        input_buffer = BytesIO(xlsx_bytes)
        wb = load_workbook(input_buffer, read_only=False, data_only=True)
        ws = wb.active
        
        # Читаем все данные в список списков для быстрого доступа
        data = list(ws.iter_rows(values_only=True))
        
        if not data:
            wb.close()
            return np.zeros((0, 0), dtype=np.float64), []
        
        # Заголовки из первой строки (начиная со 2-го столбца)
        headers = [str(cell) if cell is not None else "" for cell in data[0][1:]]
        
        # Размерность матрицы
        n = len(data) - 1  # Исключаем строку заголовков
        matrix = np.zeros((n, n), dtype=np.float64)
        
        # Заполняем матрицу
        for i in range(n):
            row_data = data[i + 1]  # +1 чтобы пропустить заголовки
            for j in range(min(n, len(row_data) - 1)):
                cell_value = row_data[j + 1]  # +1 чтобы пропустить заголовок строки
                if cell_value is not None:
                    try:
                        matrix[i, j] = float(cell_value)
                    except (ValueError, TypeError):
                        matrix[i, j] = 0.0
        
        wb.close()
        return matrix, headers

    def npy_to_matrix(self, npy_bytes: bytes) -> np.ndarray:
        """Читает numpy массив из NPY файла."""
        input_buffer = BytesIO(npy_bytes)
        return np.load(input_buffer, allow_pickle=False)

    def matrix_to_npy_bytes(self, matrix: np.ndarray) -> bytes:
        """Конвертирует numpy массив в NPY формат."""
        output = BytesIO()
        np.save(output, matrix, allow_pickle=False)
        output.seek(0)
        return output.read()

